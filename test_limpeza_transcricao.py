"""
Testes da limpeza IRaMuTeQ, importação de planilhas e transcrição
(Whisper simulado; o Whisper real é testado no GitHub Actions).
"""
from __future__ import annotations

import io
import tempfile
from contextlib import redirect_stdout
from pathlib import Path

import cli
from corpus import Document, parse_corpus
from limpeza import (
    build_header, clean_documents, clean_text, fix_header, read_table, table_to_documents,
)
from transcricao import Transcript, find_audio, load_variables, transcribe_to_documents


def test_regras_de_limpeza():
    casos = {
        'Ele disse "não" e saiu': "Ele disse não e saiu",
        "guarda-chuva e bem-te-vi": "guarda_chuva e bem_te_vi",
        "disse-me que ia fazê-lo": "disse me que ia fazê lo",
        "água d'água": "água d água",
        "Foi assim... e depois…": "Foi assim. e depois.",
        "Cresceu 50% em 2020": "Cresceu 50 por_cento em 2020",
        "arroz & feijão": "arroz e feijão",
        "nota *importante* #tag @fulano": "nota importante tag fulano",
        "veja https://exemplo.org/x agora": "veja agora",
        "ótimo 😀👍 mesmo": "ótimo mesmo",
        "isso — aquilo – outro": "isso aquilo outro",
        "  muitos\t\tespaços  ": "muitos espaços",
    }
    for entrada, esperado in casos.items():
        assert clean_text(entrada) == esperado, (entrada, clean_text(entrada))
    print("OK test_regras_de_limpeza")


def test_expressoes_compostas():
    termos = ["sistema único de saúde", "bem estar"]
    t = "O Sistema Único de Saúde garante o bem-estar e o bem estar."
    assert clean_text(t, termos) == "O sistema_único_de_saúde garante o bem_estar e o bem_estar."
    print("OK test_expressoes_compostas")


def test_linha_tematica_e_cabecalhos():
    docs = parse_corpus("**** *Suj_01 *região_Sul *idade\n-*tema_saude\nTexto \"um\".\n"
                        "**** *suj_02\n😀\n")
    limpos, rel = clean_documents(docs)
    assert len(limpos) == 1 and rel.empty_documents == [2]
    assert limpos[0].header == "**** *Suj_01 *regiao_sul"
    assert limpos[0].text == "-*tema_saude\nTexto um."
    assert any("corrigida" in p for p in rel.header_problems)
    assert any("removida" in p for p in rel.header_problems)
    assert fix_header("**** *sexo_f *idade_30") == ("**** *sexo_f *idade_30", [])
    assert build_header({"Sexo": "Feminino", "Região": "São Paulo"}) == "**** *sexo_feminino *regiao_saopaulo"
    print("OK test_linha_tematica_e_cabecalhos")


def test_importar_planilha_csv_e_xlsx():
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        csv_path = tmp / "respostas.csv"
        csv_path.write_text(
            "id;sexo;idade;resposta\n1;F;30;\"Gosto do bem-estar.\"\n2;M;45;\n3;F;22;Falta \"verba\".\n",
            encoding="cp1252",
        )
        rows = read_table(csv_path)
        docs = table_to_documents(rows, "resposta", ["sexo", "idade"])
        assert [d.header for d in docs] == ["**** *sexo_f *idade_30", "**** *sexo_f *idade_22"]
        try:
            table_to_documents(rows, "texto")
            raise AssertionError("coluna inexistente deveria falhar")
        except ValueError as e:
            assert "resposta" in str(e)

        out = tmp / "corpus.txt"
        with redirect_stdout(io.StringIO()):
            assert cli.main(["importar", str(csv_path), "--texto", "resposta",
                             "--variaveis", "sexo", "idade", "--saida", str(out)]) == 0
        assert out.read_text(encoding="utf-8") == (
            "**** *sexo_f *idade_30\nGosto do bem_estar.\n\n**** *sexo_f *idade_22\nFalta verba.\n"
        )

        try:
            from openpyxl import Workbook
        except ImportError:
            print("  (openpyxl ausente: XLSX não testado)")
        else:
            wb = Workbook()
            ws = wb.active
            ws.append(["grupo", "texto"])
            ws.append(["A", "Primeira resposta"])
            ws.append(["B", None])
            xlsx = tmp / "r.xlsx"
            wb.save(xlsx)
            docs = table_to_documents(read_table(xlsx), "texto", ["grupo"])
            assert [(d.header, d.text) for d in docs] == [("**** *grupo_a", "Primeira resposta")]
    print("OK test_importar_planilha_csv_e_xlsx")


class FakeWhisper:
    def __init__(self, textos):
        self.textos = textos

    def transcribe(self, path):
        lang, text = self.textos[path.name]
        return Transcript(path, text, lang, 90.0)


def test_transcricao_com_variaveis():
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        audios = tmp / "audios"
        audios.mkdir()
        for nome in ("Ent 01.mp3", "ent02.m4a", "ent03.wav", "notas.txt"):
            (audios / nome).write_bytes(b"x")
        (tmp / "participantes.csv").write_text(
            "arquivo,sexo,escolaridade\nEnt 01.mp3,F,Superior\nent02,M,Médio\n", encoding="utf-8"
        )
        files = find_audio([audios])
        assert [f.name for f in files] == ["Ent 01.mp3", "ent02.m4a", "ent03.wav"]

        fake = FakeWhisper({
            "Ent 01.mp3": ("pt", 'A gente espera muito no posto de saúde... "demais".'),
            "ent02.m4a": ("en", "The clinic is far away."),
            "ent03.wav": ("pt", ""),
        })
        variaveis = load_variables(tmp / "participantes.csv")
        docs, transcripts, avisos = transcribe_to_documents(files, fake, variaveis)
        assert [d.header for d in docs] == [
            "**** *ent_ent01 *sexo_f *escolaridade_superior",
            "**** *ent_ent02 *sexo_m *escolaridade_medio",
        ]
        assert any("ent03.wav" in a and "nenhuma fala" in a for a in avisos)

        out = tmp / "corpus.txt"
        buf = io.StringIO()
        args = cli.build_parser().parse_args(
            ["transcrever", str(audios), "--variaveis", str(tmp / "participantes.csv"),
             "--saida", str(out)]
        )
        with redirect_stdout(buf):
            assert cli.cmd_transcrever(args, transcriber=fake) == 0
        texto = out.read_text(encoding="utf-8")
        assert "**** *ent_ent01 *sexo_f *escolaridade_superior\nA gente espera muito no posto de saúde. demais." in texto
        assert "idioma en" in buf.getvalue() and "traduzir" in buf.getvalue()
    print("OK test_transcricao_com_variaveis")


def test_pipeline_aplica_limpeza():
    from pipeline import Config, run_pipeline
    from test_cli import FakeBridge

    with tempfile.TemporaryDirectory() as tmp:
        bridge = FakeBridge()
        cfg = Config(translate=False, use_llm=False, use_cache=False,
                     compound_terms=["posto de saúde"], cache_db=str(Path(tmp) / "c.db"))
        docs = [Document('O "posto de saúde" fecha cedo...'), Document("bem-estar"), Document("😀")]
        res = run_pipeline(docs, cfg, log=lambda m: None, bridge_factory=lambda: bridge)
        assert bridge.calls[0][0] == ["O posto_de_saúde fecha cedo.", "bem_estar"]
        assert any("vazios" in w for w in res.warnings)

        bridge2 = FakeBridge()
        run_pipeline(docs[:2], Config(translate=False, use_llm=False, use_cache=False, clean=False),
                     log=lambda m: None, bridge_factory=lambda: bridge2)
        assert bridge2.calls[0][0][0] == 'O "posto de saúde" fecha cedo...'
    print("OK test_pipeline_aplica_limpeza")


if __name__ == "__main__":
    test_regras_de_limpeza()
    test_expressoes_compostas()
    test_linha_tematica_e_cabecalhos()
    test_importar_planilha_csv_e_xlsx()
    test_transcricao_com_variaveis()
    test_pipeline_aplica_limpeza()
    print("\n=== LIMPEZA E TRANSCRIÇÃO OK ===")
