"""
Relatório em Word (.docx) pronto para o capítulo de resultados:
método, estatísticas textuais, figuras numeradas e uma seção por classe.
Para PDF: abra no Word/LibreOffice e use "Salvar como PDF".
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import TYPE_CHECKING

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

if TYPE_CHECKING:
    from pipeline import PipelineResult

INK_SECONDARY = RGBColor(0x52, 0x51, 0x4E)


def _add_rich(paragraph, text: str) -> None:
    """Converte o markdown leve dos itens de método (`código`, **negrito**)."""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
        else:
            paragraph.add_run(part)


def _caption(doc, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = INK_SECONDARY
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _table(doc, header, rows, widths_cm=None):
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Light Grid Accent 1"
    for cell, text in zip(table.rows[0].cells, header):
        cell.text = text
        cell.paragraphs[0].runs[0].bold = True
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = str(value)
    if widths_cm:
        for row in table.rows:
            for cell, w in zip(row.cells, widths_cm):
                cell.width = Cm(w)
    return table


def build_docx(result: "PipelineResult", path: str | Path) -> Path:
    from analises import FIGURE_TITLES
    from pipeline import _class_name, method_items

    path = Path(path)
    doc = DocxDocument()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    doc.add_heading(f"Relatório Textome — {result.source_name}", level=0)
    meta = doc.add_paragraph(f"Gerado em {result.finished_at}.")
    meta.runs[0].font.color.rgb = INK_SECONDARY

    doc.add_heading("Método", level=1)
    for item in method_items(result):
        _add_rich(doc.add_paragraph(style="List Bullet"), item)

    fig_n = 0
    analysis = result.analysis
    if analysis is not None:
        st = analysis.stats
        doc.add_heading("Estatísticas textuais", level=1)
        _table(doc, ["Medida", "Valor"], [
            ("Segmentos de texto", st.segments),
            ("Segmentos classificados", f"{st.classified} ({st.classified_pct:.1f}%)"),
            ("Ocorrências", st.occurrences),
            ("Formas distintas", st.forms),
            ("Hápax", f"{st.hapax} ({st.hapax_pct_forms:.1f}% das formas)"),
        ], widths_cm=[7, 6])

        doc.add_heading("Figuras", level=1)
        for key, fig_path in analysis.figures.items():
            try:
                doc.add_picture(str(fig_path), width=Cm(16))
            except Exception as e:  # uma figura corrompida não derruba o relatório
                result.warnings.append(f"Figura '{key}' não incluída no Word: {type(e).__name__}")
                continue
            fig_n += 1
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            _caption(doc, f"Figura {fig_n} – {FIGURE_TITLES.get(key, key)}. Fonte: elaborado com Textome.")

    total = sum(result.sizes.values()) or 1
    doc.add_heading("Classes", level=1)
    for cid, data in result.classes.items():
        n_seg = result.sizes.get(cid, 0)
        interp = result.interpretations.get(cid, {})
        doc.add_heading(f"Classe {cid} — {_class_name(result, cid)}", level=2)
        doc.add_paragraph(f"{n_seg} segmento(s) de texto ({100 * n_seg / total:.1f}% do corpus classificado).")
        if interp.get("descricao"):
            doc.add_paragraph().add_run(interp["descricao"]).italic = True
        if interp.get("resumo"):
            doc.add_paragraph(interp["resumo"])
        if interp.get("interpretacao"):
            doc.add_paragraph().add_run("Interpretação aprofundada").bold = True
            doc.add_paragraph(interp["interpretacao"])
        forms = data.get("forms", [])[:15]
        if forms:
            doc.add_paragraph().add_run("Formas características").bold = True
            _table(doc, ["Forma", "χ²"], [(f, f"{c:.1f}") for f, c in forms], widths_cm=[7, 3])
        segs = data.get("segments", [])[:5]
        if segs:
            doc.add_paragraph().add_run("Segmentos de texto").bold = True
            for seg in segs:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(1)
                p.add_run(seg).italic = True

    if result.warnings:
        doc.add_heading("Avisos", level=1)
        for w in result.warnings:
            doc.add_paragraph(w, style="List Bullet")

    if result.config.use_llm:
        note = doc.add_paragraph()
        run = note.add_run(
            "Nomes, descrições e resumos das classes foram sugeridos por um modelo de linguagem "
            "local e devem ser revisados pelo pesquisador antes da publicação."
        )
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = INK_SECONDARY

    doc.save(path)
    return path
