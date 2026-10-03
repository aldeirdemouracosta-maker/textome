"""
Limpeza de corpus segundo as regras usuais do IRaMuTeQ e montagem de
corpus a partir de planilhas (respostas abertas de questionário etc.).

Regras aplicadas ao texto (cabeçalhos **** são validados à parte):
- remove URLs, emojis e símbolos gráficos
- remove aspas; apóstrofo entre letras vira espaço (d'água → d água)
- une expressões compostas da lista do pesquisador (sistema único de saúde → sistema_único_de_saúde)
- hífen entre palavras vira "_" (guarda-chuva → guarda_chuva), exceto pronomes
  oblíquos, que são separados (disse-me → disse me)
- travessões viram espaço; "..." e "…" viram "."
- "%" → " por_cento"; remove * $ # @ e outros caracteres especiais; "&" → " e "
- espaços e tabulações normalizados
Linhas temáticas "-*tema" são preservadas.
"""

from __future__ import annotations

import csv
import re
import unicodedata
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

from corpus import Document

CLITICS = {
    "me", "te", "se", "lhe", "lhes", "nos", "vos", "o", "a", "os", "as",
    "lo", "la", "los", "las", "no", "na", "nas", "mo", "ma", "to", "ta",
}

_URL = re.compile(r"(https?://|www\.)\S+", re.IGNORECASE)
_QUOTES = re.compile(r"[\"“”„«»‹›]")
_APOSTROPHE = re.compile(r"(?<=\w)['’`´](?=\w)")
_LOOSE_APOSTROPHE = re.compile(r"['’`´‘]")
_HYPHENATED = re.compile(r"\b\w+(?:-\w+)+\b")
_DASH = re.compile(r"\s+[-–—]+\s+|[–—]+")
_ELLIPSIS = re.compile(r"\.{2,}|…")
_SPACES = re.compile(r"[ \t ]+")
_VAR_TOKEN = re.compile(r"^\*[A-Za-z0-9]+_[A-Za-z0-9]+$")


def _strip_accents(text: str) -> str:
    return "".join(
        c for c in unicodedata.normalize("NFD", text) if unicodedata.category(c) != "Mn"
    )


def sanitize_token(value: str) -> str:
    """Nome/modalidade de variável IRaMuTeQ: só letras sem acento e números."""
    return re.sub(r"[^A-Za-z0-9]", "", _strip_accents(str(value)).strip().lower())


def build_header(variables: Dict[str, str]) -> str:
    parts = ["****"]
    for name, value in variables.items():
        n, v = sanitize_token(name), sanitize_token(value)
        if n and v:
            parts.append(f"*{n}_{v}")
    return " ".join(parts)


def fix_header(header: str) -> Tuple[str, List[str]]:
    """Valida e corrige uma linha ****. Retorna (cabeçalho corrigido, problemas)."""
    problems: List[str] = []
    tokens = header.strip().split()
    if not tokens or tokens[0] != "****":
        problems.append("o cabeçalho deve começar com '****' seguido de espaço")
        tokens = ["****"] + [t for t in tokens if not t.startswith("****")]
    fixed = ["****"]
    for tok in tokens[1:]:
        if _VAR_TOKEN.match(tok):
            fixed.append(tok)
            continue
        name, _, value = tok.lstrip("*").partition("_")
        n, v = sanitize_token(name), sanitize_token(value)
        if n and v:
            fixed.append(f"*{n}_{v}")
            problems.append(f"variável '{tok}' corrigida para '*{n}_{v}'")
        else:
            problems.append(f"variável '{tok}' inválida (use *nome_valor) e foi removida")
    return " ".join(fixed), problems


def _compound_pattern(term: str) -> re.Pattern:
    words = [re.escape(w) for w in re.split(r"[\s_-]+", term.strip()) if w]
    return re.compile(r"(?<!\w)" + r"[\s-]+".join(words) + r"(?!\w)", re.IGNORECASE)


def _clean_line(line: str, compounds: Sequence[Tuple[re.Pattern, str]], stats: Counter) -> str:
    def sub(pattern: re.Pattern, repl, text: str, rule: str) -> str:
        new, n = pattern.subn(repl, text)
        if n:
            stats[rule] += n
        return new

    line = unicodedata.normalize("NFC", line)
    line = sub(_URL, " ", line, "URLs removidas")

    kept = []
    for ch in line:
        if unicodedata.category(ch) in ("So", "Sk", "Cs", "Co") or ord(ch) in (0xFE0F, 0x200D):
            stats["emojis/símbolos removidos"] += 1
        else:
            kept.append(ch)
    line = "".join(kept)

    line = sub(_QUOTES, "", line, "aspas removidas")
    line = sub(_APOSTROPHE, " ", line, "apóstrofos trocados por espaço")
    line = sub(_LOOSE_APOSTROPHE, "", line, "aspas removidas")

    for pattern, joined in compounds:
        line = sub(pattern, joined, line, "expressões compostas unidas")

    def hyphen(match: re.Match) -> str:
        parts = match.group(0).split("-")
        if len(parts) == 2 and parts[1].lower() in CLITICS:  # disse-me, fazê-lo
            stats["hífens de pronome separados"] += 1
            return " ".join(parts)
        stats["hífens trocados por _"] += 1  # guarda-chuva, bem-te-vi
        return "_".join(parts)

    line = _HYPHENATED.sub(hyphen, line)

    line = sub(_DASH, " ", line, "travessões removidos")
    line = sub(_ELLIPSIS, ".", line, "reticências normalizadas")
    line = sub(re.compile(r"\s*%"), " por_cento", line, "% trocado por por_cento")
    line = sub(re.compile(r"\s*&\s*"), " e ", line, "& trocado por e")
    line = sub(re.compile(r"[*$#@^~|\\{}\[\]<>=+]"), " ", line, "caracteres especiais removidos")
    line = _SPACES.sub(" ", line).strip()
    return line


@dataclass
class CleaningReport:
    changes: Counter = field(default_factory=Counter)
    header_problems: List[str] = field(default_factory=list)
    empty_documents: List[int] = field(default_factory=list)

    def summary_lines(self) -> List[str]:
        lines = [f"{rule}: {n}" for rule, n in self.changes.most_common()]
        lines += self.header_problems
        if self.empty_documents:
            lines.append(f"documentos vazios removidos: {self.empty_documents}")
        return lines


def clean_text(text: str, compound_terms: Iterable[str] = (), stats: Optional[Counter] = None) -> str:
    stats = stats if stats is not None else Counter()
    compounds = [
        (_compound_pattern(t), re.sub(r"[\s-]+", "_", t.strip()))
        for t in compound_terms if t and t.strip()
    ]
    out = []
    for line in text.splitlines():
        if line.strip().startswith("-*"):  # linha temática do IRaMuTeQ
            out.append(line.strip())
            continue
        cleaned = _clean_line(line, compounds, stats)
        if cleaned:
            out.append(cleaned)
    return "\n".join(out)


def clean_documents(
    docs: List[Document], compound_terms: Iterable[str] = ()
) -> Tuple[List[Document], CleaningReport]:
    report = CleaningReport()
    terms = list(compound_terms)
    cleaned: List[Document] = []
    for i, doc in enumerate(docs, 1):
        header = doc.header
        if header:
            header, problems = fix_header(header)
            report.header_problems += [f"documento {i}: {p}" for p in problems]
        text = clean_text(doc.text, terms, report.changes)
        if not text.strip():
            report.empty_documents.append(i)
            continue
        cleaned.append(Document(text, header))
    return cleaned, report


def load_compound_terms(path: str | Path) -> List[str]:
    """Um termo por linha; linhas vazias e iniciadas por # são ignoradas."""
    lines = Path(path).read_text(encoding="utf-8").splitlines()
    return [ln.strip() for ln in lines if ln.strip() and not ln.strip().startswith("#")]


# ---------------------------------------------------------------------------
# Planilhas → corpus
# ---------------------------------------------------------------------------

def read_table(path: str | Path) -> List[Dict[str, str]]:
    """Lê CSV (separador detectado: , ; ou tab) ou XLSX (requer openpyxl)."""
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError as e:
            raise ImportError("Para ler .xlsx instale: pip install openpyxl") from e
        ws = load_workbook(path, read_only=True, data_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        header = [str(h).strip() if h is not None else "" for h in rows[0]]
        return [
            {h: ("" if v is None else str(v)) for h, v in zip(header, row) if h}
            for row in rows[1:]
        ]

    from corpus import decode_bytes

    raw = decode_bytes(path.read_bytes())
    dialect = csv.Sniffer().sniff(raw[:5000], delimiters=",;\t")
    return [
        {(k or "").strip(): (v or "") for k, v in row.items()}
        for row in csv.DictReader(raw.splitlines(), dialect=dialect)
    ]


def table_to_documents(
    rows: List[Dict[str, str]], text_column: str, variable_columns: Sequence[str] = ()
) -> List[Document]:
    if rows and text_column not in rows[0]:
        raise ValueError(
            f"Coluna de texto '{text_column}' não encontrada. Colunas: {', '.join(rows[0])}"
        )
    missing = [c for c in variable_columns if rows and c not in rows[0]]
    if missing:
        raise ValueError(f"Colunas de variáveis não encontradas: {', '.join(missing)}")
    docs = []
    for row in rows:
        text = (row.get(text_column) or "").strip()
        if text:
            docs.append(Document(text, build_header({c: row.get(c, "") for c in variable_columns})))
    return docs
