"""
Testes das análises (AFC, similitude, nuvem, estatísticas) e do relatório Word.
"""
from __future__ import annotations

import io
import re
import tempfile
from collections import Counter
from contextlib import redirect_stdout
from pathlib import Path

import numpy as np
from scipy import sparse

from analises import (
    CorpusMatrix, class_term_table, cooccurrence_graph, correspondence_analysis,
    run_all, term_classes, text_stats,
)
from corpus import parse_corpus


def matrix_from_texts(texts, groups):
    toks = [re.findall(r"\w+", t.lower()) for t in texts]
    vocab = sorted({w for t in toks for w in t})
    idx = {w: i for i, w in enumerate(vocab)}
    r, c, v = [], [], []
    for i, t in enumerate(toks):
        for w in t:
            r.append(i); c.append(idx[w]); v.append(1)
    return CorpusMatrix(sparse.csr_matrix((v, (r, c)), shape=(len(texts), len(vocab))), vocab, groups)


def demo_matrix():
    docs = parse_corpus(Path(__file__).with_name("exemplos").joinpath("corpus_demo.txt")
                        .read_text(encoding="utf-8"))
    temas = [re.search(r"\*tema_(\w+)", d.header).group(1) for d in docs]
    ids = {t: i + 1 for i, t in enumerate(sorted(set(temas)))}
    return matrix_from_texts([d.text for d in docs], [ids[t] for t in temas]), ids


def test_estatisticas_textuais():
    m = matrix_from_texts(["a b b c", "c d", "e"], [1, 2, 0])
    st = text_stats(m)
    assert (st.segments, st.classified, st.occurrences, st.forms, st.hapax) == (3, 2, 7, 5, 3)
    assert round(st.hapax_pct_forms) == 60
    print("OK test_estatisticas_textuais")


def test_afc_matematica():
    N = np.array([[10, 2, 0, 5], [1, 8, 4, 0], [0, 3, 9, 2]], dtype=float)
    rows, cols, pct = correspondence_analysis(N, 2)
    total = N.sum()
    r, c = N.sum(1) / total, N.sum(0) / total
    # centróides ponderados na origem
    assert np.allclose(r @ rows, 0) and np.allclose(c @ cols, 0)
    # inércia total = χ² / n
    expected = np.outer(N.sum(1), N.sum(0)) / total
    chi2 = ((N - expected) ** 2 / expected).sum()
    eig_sum = (r[:, None] * rows ** 2).sum()  # inércia nos 2 fatores = toda (3 linhas → 2 fatores)
    assert np.isclose(eig_sum, chi2 / total)
    assert np.isclose(pct.sum(), 100)
    print("OK test_afc_matematica")


def test_tabela_e_classes_das_formas():
    m = matrix_from_texts(["saude hospital", "saude fila", "escola professor", "escola aula", "x"],
                          [1, 1, 2, 2, 0])
    ids, table = class_term_table(m)
    assert ids == [1, 2] and table.sum() == 8, "segmento sem classe fica fora"
    tc = term_classes(m)
    assert tc["saude"] == 1 and tc["escola"] == 2
    tc2 = term_classes(m, {2: [("saude", 99.0)]})
    assert tc2["saude"] == 2, "χ² do rainette tem prioridade"
    print("OK test_tabela_e_classes_das_formas")


def test_grafo_de_similitude():
    m = matrix_from_texts(["a b c", "a b", "a c", "b c d", "e f"], [1, 1, 2, 2, 3])
    g, tree = cooccurrence_graph(m, top_n=10, min_cooc=1)
    assert g["a"]["b"]["weight"] == 2
    assert tree.number_of_edges() == tree.number_of_nodes() - nx_components(tree)
    print("OK test_grafo_de_similitude")


def nx_components(g):
    import networkx as nx
    return nx.number_connected_components(g)


def test_figuras_do_corpus_exemplo():
    m, ids = demo_matrix()
    names = {v: k for k, v in ids.items()}
    with tempfile.TemporaryDirectory() as tmp:
        out = run_all(m, Path(tmp), dict(Counter(m.groups)), names, dendrogram_png=b"\x89PNG fake")
        assert set(out.figures) == {"dendrograma", "classes", "afc", "similitude", "nuvem"}
        for p in out.figures.values():
            assert p.exists() and p.stat().st_size > 1000 or p.name == "dendrograma.png"
        assert out.ca is not None and len(out.ca.inertia_pct) == 2
    print("OK test_figuras_do_corpus_exemplo")


def test_afc_com_duas_classes_avisa():
    m = matrix_from_texts(["a b", "c d", "a c"], [1, 2, 1])
    with tempfile.TemporaryDirectory() as tmp:
        out = run_all(m, Path(tmp), {1: 2, 2: 1}, {})
        assert "afc" not in out.figures and any("3 classes" in w for w in out.warnings)
    print("OK test_afc_com_duas_classes_avisa")


class BridgeComMatriz:
    """FakeBridge que também fornece a matriz e o dendrograma."""

    def __init__(self, dendro_error=None):
        self.dendro_error = dendro_error

    def run(self, texts, **kwargs):
        self.texts = texts
        groups = [1 + (i % 3) for i in range(len(texts))]
        self.groups = groups
        return {cid: {"forms": [], "segments": [t for t, g in zip(texts, groups) if g == cid][:3]}
                for cid in (1, 2, 3)}

    def get_group_sizes(self):
        return dict(Counter(self.groups))

    def get_matrix(self):
        return matrix_from_texts(self.texts, self.groups)

    def save_dendrogram(self, path):
        if self.dendro_error == "png corrompido":
            Path(path).write_bytes(b"\x89PNG\r\n\x1a\n" + b"0" * 2000)
            return None
        if self.dendro_error:
            return self.dendro_error
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        fig, ax = plt.subplots(figsize=(4, 3))
        ax.plot([0, 1], [0, 1])
        fig.savefig(path)
        plt.close(fig)
        return None


def test_pipeline_gera_figuras_e_word():
    from docx import Document as Docx

    import cli

    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        conf = tmp / "p.yaml"
        conf.write_text(f"cache_db: {tmp / 'c.db'}\n", encoding="utf-8")
        args = cli.build_parser().parse_args(
            ["run", "exemplos/corpus_demo.txt", "--config", str(conf), "--mock", "--saida", str(tmp / "out")]
        )
        with redirect_stdout(io.StringIO()):
            assert cli.cmd_run(args, bridge_factory=lambda: BridgeComMatriz(dendro_error="sem R")) == 0
        out = tmp / "out"
        figs = sorted(p.name for p in (out / "figuras").iterdir())
        assert figs == ["afc.png", "classes.png", "nuvem.png", "similitude.png"], figs

        md = (out / "relatorio.md").read_text(encoding="utf-8")
        assert "## Estatísticas textuais" in md and "](figuras/afc.png)" in md
        assert "Dendrograma não gerado: sem R" in md

        docx = Docx(str(out / "relatorio.docx"))
        text = "\n".join(p.text for p in docx.paragraphs)
        assert "Método" in text and "Figura 1 – Distribuição" in text and "Classe 1" in text
        assert "análise fatorial de correspondência" in text
        assert len(docx.inline_shapes) == 4
        assert any(t.cell(0, 0).text == "Medida" for t in docx.tables)
        # Dendrograma válido entra como Figura 1; um PNG corrompido vira aviso, sem derrubar o Word
        for erro, n_figs, aviso in ((None, 5, None), ("png corrompido", 4, "não incluída no Word")):
            saida = tmp / f"out_{n_figs}"
            args = cli.build_parser().parse_args(
                ["run", "exemplos/corpus_demo.txt", "--config", str(conf), "--mock", "--saida", str(saida)]
            )
            with redirect_stdout(io.StringIO()):
                assert cli.cmd_run(args, bridge_factory=lambda e=erro: BridgeComMatriz(dendro_error=e)) == 0
            d = Docx(str(saida / "relatorio.docx"))
            assert len(d.inline_shapes) == n_figs
            texto = "\n".join(p.text for p in d.paragraphs)
            if aviso:
                assert aviso in texto
            else:
                assert "Figura 1 – Dendrograma" in texto
    print("OK test_pipeline_gera_figuras_e_word")


if __name__ == "__main__":
    test_estatisticas_textuais()
    test_afc_matematica()
    test_tabela_e_classes_das_formas()
    test_grafo_de_similitude()
    test_figuras_do_corpus_exemplo()
    test_afc_com_duas_classes_avisa()
    test_pipeline_gera_figuras_e_word()
    print("\n=== ANÁLISES E RELATÓRIO OK ===")
